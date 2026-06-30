import copy
from docx.oxml.ns import qn

def get_tr_list(table):
    return table._tbl.findall(qn('w:tr'))

def get_tc(tr, col_idx):
    """Get the col_idx-th <w:tc> direct child of a <w:tr>, ignoring python-docx merge resolution."""
    return tr.findall(qn('w:tc'))[col_idx]

def set_tc_text(tc, text, bold=False, font_size=24, align=None):
    for p in tc.findall(qn('w:p')):
        tc.remove(p)
    lines = text.split('\n') if text else ['']
    for line in lines:
        p = tc.makeelement(qn('w:p'), {})
        pPr = p.makeelement(qn('w:pPr'), {})
        if align:
            jc = p.makeelement(qn('w:jc'), {qn('w:val'): align})
            pPr.append(jc)
        spacing = p.makeelement(qn('w:spacing'), {qn('w:line'): '276', qn('w:lineRule'): 'auto'})
        pPr.append(spacing)
        p.append(pPr)
        if line != '':
            r = p.makeelement(qn('w:r'), {})
            rPr = r.makeelement(qn('w:rPr'), {})
            rFonts = r.makeelement(qn('w:rFonts'), {
                qn('w:ascii'): 'Times New Roman', qn('w:eastAsia'): '標楷體',
                qn('w:hAnsi'): 'Times New Roman', qn('w:cs'): 'Times New Roman'})
            rPr.append(rFonts)
            if bold:
                rPr.append(r.makeelement(qn('w:b'), {}))
            rPr.append(r.makeelement(qn('w:sz'), {qn('w:val'): str(font_size)}))
            rPr.append(r.makeelement(qn('w:szCs'), {qn('w:val'): str(font_size)}))
            r.append(rPr)
            t = r.makeelement(qn('w:t'), {qn('xml:space'): 'preserve'})
            t.text = line
            r.append(t)
            p.append(r)
        tc.append(p)

def set_tc_vmerge(tc, val):
    """val: 'restart', 'continue' (represented as no w:val attr), or None to remove vMerge entirely."""
    tcPr = tc.find(qn('w:tcPr'))
    if tcPr is None:
        tcPr = tc.makeelement(qn('w:tcPr'), {})
        tc.insert(0, tcPr)
    vMerge = tcPr.find(qn('w:vMerge'))
    if val is None:
        if vMerge is not None:
            tcPr.remove(vMerge)
        return
    if vMerge is None:
        vMerge = tcPr.makeelement(qn('w:vMerge'), {})
        tcPr.insert(0, vMerge)
    if val == 'continue':
        if qn('w:val') in vMerge.attrib:
            del vMerge.attrib[qn('w:val')]
    else:
        vMerge.set(qn('w:val'), val)

def set_tc_shade(tc, fill):
    tcPr = tc.find(qn('w:tcPr'))
    if tcPr is None:
        tcPr = tc.makeelement(qn('w:tcPr'), {})
        tc.insert(0, tcPr)
    shd = tcPr.find(qn('w:shd'))
    if shd is not None:
        tcPr.remove(shd)
    shd = tcPr.makeelement(qn('w:shd'), {qn('w:val'): 'clear', qn('w:color'): 'auto', qn('w:fill'): fill})
    tcPr.append(shd)

def duplicate_tr(table, row_index):
    tbl = table._tbl
    trs = tbl.findall(qn('w:tr'))
    src = trs[row_index]
    new_tr = copy.deepcopy(src)
    src.addnext(new_tr)
    return new_tr

def remove_tr(table, row_index):
    tbl = table._tbl
    trs = tbl.findall(qn('w:tr'))
    tbl.remove(trs[row_index])

def highlight_table6(table, selected_subitems):
    """
    Shade selected 學習表現子項 rows in table6 (學習表現三項目十一子項) with yellow
    highlight (FFE699). Also shades the 項目 label cell (col0) for any group that
    contains at least one selected sub-item.

    Do NOT call table.cell() anywhere in this function — col0 has vMerge chains
    across each 項目 group, and table.cell() would silently resolve all rows in a
    group to the same XML element.

    Args:
        table: the python-docx Table object for doc.tables[6]
        selected_subitems: set or list of 學習表現子項 text strings to highlight,
            e.g. {"推理論證", "建立模型", "觀察與定題", "計畫與執行",
                  "培養科學探究的興趣"}
            Derive this from the 子項 column of table7 (學習表現檢核點).

    Row index → 學習表現子項 mapping (0 = header row, do not touch):
        1=想像創造, 2=推理論證, 3=批判思辨, 4=建立模型,
        5=觀察與定題, 6=計畫與執行, 7=分析與發現, 8=討論與傳達,
        9=培養科學探究的興趣, 10=養成應用科學思考與探究的習慣, 11=認識科學的本質

    col0 vMerge restart rows (項目 label cells):
        1 = 探究能力-思考智能 (spans rows 1-4)
        5 = 探究能力-問題解決 (spans rows 5-8)
        9 = 科學的態度與本質  (spans rows 9-11)
    """
    HIGHLIGHT = 'FFE699'
    trs = get_tr_list(table)

    subitem_rows = {
        '想像創造': 1, '推理論證': 2, '批判思辨': 3, '建立模型': 4,
        '觀察與定題': 5, '計畫與執行': 6, '分析與發現': 7, '討論與傳達': 8,
        '培養科學探究的興趣': 9,
        '養成應用科學思考與探究的習慣': 10,
        '認識科學的本質': 11,
    }
    # col0 label row index → set of row indices in that group
    group_label_to_members = {
        1: frozenset([1, 2, 3, 4]),
        5: frozenset([5, 6, 7, 8]),
        9: frozenset([9, 10, 11]),
    }

    selected_set = set(selected_subitems)
    selected_row_indices = {subitem_rows[s] for s in selected_set if s in subitem_rows}

    # shade col1 (子項) for each selected row
    for ridx in selected_row_indices:
        set_tc_shade(get_tc(trs[ridx], 1), HIGHLIGHT)

    # shade col0 (項目 label) if ANY member of the group is selected
    for label_ridx, members in group_label_to_members.items():
        if members & selected_row_indices:
            set_tc_shade(get_tc(trs[label_ridx], 0), HIGHLIGHT)
